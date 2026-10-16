"""
=============================================================================
MÓDULO DE EXPORTACIÓN A EXCEL (openpyxl)
=============================================================================
Este módulo gestiona la generación de reportes y hojas de cálculo en formato
Microsoft Excel (.xlsx) para el sistema de inventario y facturación.

Funcionalidades principales:
1. Organización jerárquica de reportes de ventas:
   - Archivos agrupados por Año y Mes ('YYYY-MM.xlsx').
   - Hojas independientes para cada día ('Día 01', 'Día 02', etc.).
   - Actualización inteligente: si se vuelve a emitir o exportar en el mismo día,
     se reemplaza únicamente la hoja correspondiente sin alterar los demás días.
   - Detalle desglosado de facturas emitidas y productos individuales vendidos.
2. Exportación de inventario general de productos con estado de stock e indicadores (KPIs).
3. Exportación del directorio completo de clientes.
4. Aplicación de formato visual profesional:
   - Paleta corporativa sobria (azul marino oscuro / grisáceo).
   - Ajuste automático de anchos de columna (auto-width).
   - Formato numérico de moneda (USD '$' y Bolívares 'Bs.').
   - Bordes delgados y dobles para filas de totales.
=============================================================================
"""

import os
import sys
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import database


def get_reports_folder():
    """
    Obtiene la ruta absoluta al directorio de almacenamiento de reportes ('reportes').
    
    Crea la carpeta automáticamente en caso de no existir.
    Soporta compatibilidad tanto en modo desarrollo (directorio del script)
    como en ejecutables empaquetados mediante PyInstaller (sys.executable).

    Returns:
        str: Ruta absoluta a la carpeta 'reportes'.
    """
    # Determinar ruta base según si el programa corre compilado o como script
    if getattr(sys, 'frozen', False):
        base_dir = os.path.dirname(sys.executable)
    else:
        base_dir = os.path.dirname(os.path.abspath(__file__))
    
    # Construir ruta a la subcarpeta 'reportes' y garantizar su existencia
    reports_dir = os.path.join(base_dir, "reportes")
    os.makedirs(reports_dir, exist_ok=True)
    return reports_dir


def open_reports_folder():
    """
    Abre la carpeta 'reportes' en el explorador de archivos nativo de Windows.
    Permite al usuario acceder de inmediato a los archivos Excel generados.
    """
    folder = get_reports_folder()
    if os.path.exists(folder):
        os.startfile(folder)


def create_header_style():
    """
    Construye y retorna los estilos estandarizados de encabezado de tabla.
    
    Aplica fondo azul marino oscuro (#1E1E2E), texto blanco en negrita (Calibri 11)
    y alineación centrada tanto vertical como horizontalmente con ajuste de texto.

    Returns:
        tuple: (PatternFill, Font, Alignment) listos para aplicar a celdas openpyxl.
    """
    fill = PatternFill(start_color="1E1E2E", end_color="1E1E2E", fill_type="solid")
    font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return fill, font, align


def apply_auto_width(ws):
    """
    Calcula dinámicamente y asigna el ancho óptimo de cada columna de la hoja de cálculo.
    
    Recorre todas las celdas de cada columna, mide la longitud del contenido en texto
    y establece el ancho con un margen de holgura (+4 caracteres), asegurando un
    ancho mínimo de 12 unidades para evitar recortes visuales.

    Args:
        ws (openpyxl.worksheet.worksheet.Worksheet): Hoja de cálculo a ajustar.
    """
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        # Asignar ancho con holgura y mínimo de 12
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def export_invoices_by_day_and_month(invoices=None):
    """
    Exporta las facturas en la carpeta 'reportes' organizadas por libro mensual y hojas diarias.
    
    Estructura jerárquica:
    - Archivo: 'YYYY-MM.xlsx' (un archivo por cada mes/año).
    - Hojas: 'Día 01', 'Día 02', ..., 'Día 31'.
    
    Comportamiento:
    - Si ya existe un reporte previo para un día específico, sobreescribe/reemplaza
      dicha hoja con la versión más reciente sin afectar los demás días del mes.
    - Cada hoja diaria incluye:
      1. Título corporativo con nombre del negocio y fecha completa.
      2. KPIs del día (total facturas, ventas en USD y ventas en Bs).
      3. Tabla detallada de facturas emitidas (datos del cliente, método de pago, tasa, IVA, totales).
      4. Fila de resumen de totales diarios con borde contable doble.
      5. Tabla desglosada de productos vendidos durante la jornada.
      6. Totales globales de unidades y montos de productos.

    Args:
        invoices (list, opcional): Lista de diccionarios de facturas.
                                   Si es None, se consultan todas las facturas desde database.

    Returns:
        list: Lista de rutas absolutas de archivos Excel generados o actualizados.
    """
    reports_dir = get_reports_folder()
    if invoices is None:
        invoices = database.get_all_invoices()

    # Si no hay facturas registradas, retornar lista vacía
    if not invoices:
        return []

    # 1. Agrupar facturas por Año-Mes (ej. '2026-10') y Día (ej. '12')
    grouped = {}
    for inv in invoices:
        date_str = str(inv['fecha'])
        ym = date_str[:7]     # 'YYYY-MM'
        day = date_str[8:10]  # 'DD'
        if ym not in grouped:
            grouped[ym] = {}
        if day not in grouped[ym]:
            grouped[ym][day] = []
        grouped[ym][day].append(inv)

    saved_files = []
    
    # 2. Configurar estilos y bordes compartidos
    fill_header, font_header, align_header = create_header_style()
    thin_border_side = Side(style='thin', color='CBD5E1')
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    double_bottom_border = Border(top=Side(style='thin', color='94A3B8'), bottom=Side(style='double', color='1E1E2E'))

    # 3. Procesar cada archivo mensual
    for ym, days_dict in grouped.items():
        filename = f"{ym}.xlsx"
        filepath = os.path.join(reports_dir, filename)

        # Cargar libro existente si existe; si no, crear uno nuevo
        if os.path.exists(filepath):
            try:
                wb = openpyxl.load_workbook(filepath)
            except Exception:
                wb = openpyxl.Workbook()
        else:
            wb = openpyxl.Workbook()

        # 4. Procesar cada día dentro del mes
        for day in sorted(days_dict.keys()):
            sheet_title = f"Día {day}"

            # Si ya existía una hoja para este día, se elimina para recrearla con datos frescos
            if sheet_title in wb.sheetnames:
                del wb[sheet_title]

            ws = wb.create_sheet(title=sheet_title)
            day_invoices = days_dict[day]
            full_date = f"{ym}-{day}"

            # ---------------------------------------------------------
            # ENCABEZADO PRINCIPAL Y DATOS DEL NEGOCIO
            # ---------------------------------------------------------
            negocio = database.get_negocio_config()
            ws.merge_cells("A1:J1")
            title_cell = ws["A1"]
            title_cell.value = f"📊 REPORTE DE VENTAS - {negocio['nombre'].upper()} - {full_date} (DÍA {day})"
            title_cell.font = Font(name="Calibri", size=15, bold=True, color="1E1E2E")
            title_cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[1].height = 28

            # Métricas rápidas (KPIs) del día
            total_facturas = len(day_invoices)
            total_dolar_dia = sum(inv['total_dolar'] for inv in day_invoices)
            total_bs_dia = sum(inv['total_bs'] for inv in day_invoices)
            total_iva_dia = sum(inv['iva'] for inv in day_invoices)

            ws["A2"] = "Facturas Emitidas:"
            ws["A2"].font = Font(name="Calibri", size=10, bold=True, color="64748B")
            ws["B2"] = total_facturas
            ws["B2"].font = Font(name="Calibri", size=11, bold=True)

            ws["D2"] = "Total en Dólares ($):"
            ws["D2"].font = Font(name="Calibri", size=10, bold=True, color="64748B")
            ws["E2"] = total_dolar_dia
            ws["E2"].font = Font(name="Calibri", size=11, bold=True, color="059669")
            ws["E2"].number_format = '$#,##0.00'

            ws["G2"] = "Total en Bolívares (Bs):"
            ws["G2"].font = Font(name="Calibri", size=10, bold=True, color="64748B")
            ws["H2"] = total_bs_dia
            ws["H2"].font = Font(name="Calibri", size=11, bold=True, color="2563EB")
            ws["H2"].number_format = r'"Bs. "#,##0.00'

            # ---------------------------------------------------------
            # SECCIÓN 1: FACTURAS EMITIDAS EN EL DÍA
            # ---------------------------------------------------------
            ws.merge_cells("A4:J4")
            sec1_cell = ws["A4"]
            sec1_cell.value = "📋 FACTURAS EMITIDAS"
            sec1_cell.font = Font(name="Calibri", size=12, bold=True, color="1E1E2E")
            sec1_cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
            sec1_cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[4].height = 24

            headers_fac = [
                "Nº Factura", "Hora", "Cédula Cliente", "Nombre Cliente", 
                "Teléfono", "Método de Pago", "Tasa (Bs/$)", "IVA ($)", "Total ($)", "Total (Bs)"
            ]
            ws.row_dimensions[5].height = 24
            for col_idx, h in enumerate(headers_fac, start=1):
                c = ws.cell(row=5, column=col_idx, value=h)
                c.fill = fill_header
                c.font = font_header
                c.alignment = align_header
                c.border = cell_border

            # Poblar filas de facturas
            curr_row = 6
            for inv in day_invoices:
                fecha_str = str(inv['fecha'])
                hora_str = fecha_str[11:19] if len(fecha_str) >= 19 else fecha_str

                row_vals = [
                    f"FAC-{inv['id']:04d}",
                    hora_str,
                    inv['cedula_cliente'],
                    inv['cliente_nombre'] or "N/A",
                    inv['cliente_telefono'] or "N/A",
                    inv['metodo_pago'],
                    inv['tasa'],
                    inv['iva'],
                    inv['total_dolar'],
                    inv['total_bs']
                ]

                ws.row_dimensions[curr_row].height = 20
                for col_idx, val in enumerate(row_vals, start=1):
                    cell = ws.cell(row=curr_row, column=col_idx, value=val)
                    cell.border = cell_border
                    
                    # Alineación y formato específico por columna
                    if col_idx in (1, 2, 3, 5, 6):
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif col_idx in (7, 8, 9, 10):
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        if col_idx == 7:
                            cell.number_format = '#,##0.00'
                        elif col_idx in (8, 9):
                            cell.number_format = '$#,##0.00'
                        elif col_idx == 10:
                            cell.number_format = r'"Bs. "#,##0.00'
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")

                curr_row += 1

            # Fila de totales de facturación diaria
            ws.row_dimensions[curr_row].height = 22
            tot_fac_cell = ws.cell(row=curr_row, column=1, value="TOTAL DEL DÍA")
            tot_fac_cell.font = Font(name="Calibri", size=11, bold=True)
            tot_fac_cell.fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
            tot_fac_cell.border = double_bottom_border

            for c_idx in range(2, 11):
                cell = ws.cell(row=curr_row, column=c_idx)
                cell.fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
                cell.border = double_bottom_border

            # Total IVA
            c_iva = ws.cell(row=curr_row, column=8, value=total_iva_dia)
            c_iva.font = Font(name="Calibri", size=11, bold=True)
            c_iva.number_format = '$#,##0.00'
            c_iva.alignment = Alignment(horizontal="right", vertical="center")

            # Total USD
            c_usd = ws.cell(row=curr_row, column=9, value=total_dolar_dia)
            c_usd.font = Font(name="Calibri", size=11, bold=True, color="059669")
            c_usd.number_format = '$#,##0.00'
            c_usd.alignment = Alignment(horizontal="right", vertical="center")

            # Total Bs
            c_bs = ws.cell(row=curr_row, column=10, value=total_bs_dia)
            c_bs.font = Font(name="Calibri", size=11, bold=True, color="2563EB")
            c_bs.number_format = r'"Bs. "#,##0.00'
            c_bs.alignment = Alignment(horizontal="right", vertical="center")

            curr_row += 2

            # ---------------------------------------------------------
            # SECCIÓN 2: DETALLE DE PRODUCTOS INDIVIDUALES VENDIDOS
            # ---------------------------------------------------------
            ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=6)
            sec2_cell = ws.cell(row=curr_row, column=1, value="📦 DETALLE DE PRODUCTOS VENDIDOS EN EL DÍA")
            sec2_cell.font = Font(name="Calibri", size=12, bold=True, color="1E1E2E")
            sec2_cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
            sec2_cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[curr_row].height = 24
            curr_row += 1

            headers_prod = [
                "Nº Factura", "ID Producto", "Nombre del Producto", 
                "Cantidad", "Precio Unitario ($)", "Subtotal ($)"
            ]
            ws.row_dimensions[curr_row].height = 24
            for col_idx, h in enumerate(headers_prod, start=1):
                c = ws.cell(row=curr_row, column=col_idx, value=h)
                c.fill = fill_header
                c.font = font_header
                c.alignment = align_header
                c.border = cell_border
            curr_row += 1

            total_items_qty = 0
            total_items_subtotal = 0.0

            # Iterar cada factura y consultar los productos vendidos
            for inv in day_invoices:
                _, items = database.get_invoice_details(inv['id'])
                for item in items:
                    unit_p = item['producto_precio']
                    subtot = round(item['cantidad'] * unit_p, 2)
                    total_items_qty += item['cantidad']
                    total_items_subtotal += subtot

                    p_vals = [
                        f"FAC-{inv['id']:04d}",
                        item['id_producto'],
                        item['producto_nombre'],
                        item['cantidad'],
                        unit_p,
                        subtot
                    ]

                    ws.row_dimensions[curr_row].height = 20
                    for col_idx, val in enumerate(p_vals, start=1):
                        cell = ws.cell(row=curr_row, column=col_idx, value=val)
                        cell.border = cell_border
                        if col_idx in (1, 2, 4):
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                        elif col_idx in (5, 6):
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                            cell.number_format = '$#,##0.00'
                        else:
                            cell.alignment = Alignment(horizontal="left", vertical="center")

                    curr_row += 1

            # Totales de productos vendidos
            ws.row_dimensions[curr_row].height = 22
            tot_prod_cell = ws.cell(row=curr_row, column=1, value="TOTAL PRODUCTOS VENDIDOS")
            tot_prod_cell.font = Font(name="Calibri", size=11, bold=True)
            tot_prod_cell.fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
            tot_prod_cell.border = double_bottom_border

            for c_idx in range(2, 7):
                cell = ws.cell(row=curr_row, column=c_idx)
                cell.fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
                cell.border = double_bottom_border

            tot_qty_cell = ws.cell(row=curr_row, column=4, value=total_items_qty)
            tot_qty_cell.font = Font(name="Calibri", size=11, bold=True)
            tot_qty_cell.alignment = Alignment(horizontal="center", vertical="center")

            tot_sub_cell = ws.cell(row=curr_row, column=6, value=total_items_subtotal)
            tot_sub_cell.font = Font(name="Calibri", size=11, bold=True, color="059669")
            tot_sub_cell.number_format = '$#,##0.00'
            tot_sub_cell.alignment = Alignment(horizontal="right", vertical="center")

            # Aplicar auto-width a las columnas de la hoja
            apply_auto_width(ws)

        # Eliminar la pestaña en blanco por defecto 'Sheet' si existen otras hojas
        if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
            del wb["Sheet"]

        # Ordenar las pestañas cronológicamente ('Día 01', 'Día 02', etc.)
        wb._sheets.sort(key=lambda s: s.title)

        wb.save(filepath)
        saved_files.append(filepath)

    return saved_files


def export_products_to_excel(filepath=None):
    """
    Exporta el catálogo completo de productos del inventario a un archivo Excel profesional.
    
    Incluye:
    - Encabezado con información del negocio y fecha/hora de generación.
    - Indicadores clave (KPIs): Total de productos, Unidades totales en stock y Valor monetario total.
    - Tabla detallada con estado visual de stock (alerta roja si stock <= 5, verde si normal).
    - Totales generales con formato contable.

    Args:
        filepath (str, opcional): Ruta de destino del archivo. Si es None, se guarda como
                                  'reportes/Inventario_Productos.xlsx'.

    Returns:
        str: Ruta absoluta del archivo Excel generado.
    """
    if filepath is None:
        reports_dir = get_reports_folder()
        filepath = os.path.join(reports_dir, "Inventario_Productos.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario de Productos"

    fill_header, font_header, align_header = create_header_style()
    thin_border_side = Side(style='thin', color='CBD5E1')
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    double_bottom_border = Border(top=Side(style='thin', color='94A3B8'), bottom=Side(style='double', color='1E1E2E'))

    # 1. Título corporativo
    negocio = database.get_negocio_config()
    ws.merge_cells("A1:F1")
    title_cell = ws["A1"]
    title_cell.value = f"📦 INVENTARIO GENERAL - {negocio['nombre'].upper()} ({datetime.now().strftime('%d/%m/%Y %H:%M')})"
    title_cell.font = Font(name="Calibri", size=15, bold=True, color="1E1E2E")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    products = database.get_all_products()

    # 2. Métricas clave (KPIs)
    total_prods = len(products)
    total_stock = sum(p['cant'] for p in products)
    total_val_stock = sum(p['cant'] * p['precio'] for p in products)

    ws["A2"] = "Total Productos:"
    ws["A2"].font = Font(name="Calibri", size=10, bold=True, color="64748B")
    ws["B2"] = total_prods
    ws["B2"].font = Font(name="Calibri", size=11, bold=True)

    ws["C2"] = "Unidades en Stock:"
    ws["C2"].font = Font(name="Calibri", size=10, bold=True, color="64748B")
    ws["D2"] = total_stock
    ws["D2"].font = Font(name="Calibri", size=11, bold=True)

    ws["E2"] = "Valor Total Stock ($):"
    ws["E2"].font = Font(name="Calibri", size=10, bold=True, color="64748B")
    ws["F2"] = total_val_stock
    ws["F2"].font = Font(name="Calibri", size=11, bold=True, color="059669")
    ws["F2"].number_format = '$#,##0.00'

    # 3. Cabeceras de la tabla
    headers = [
        "ID Producto", "Nombre del Producto", "Cantidad (Stock)", 
        "Estado de Stock", "Precio Unitario ($)", "Valor Total Stock ($)"
    ]
    header_row = 4
    ws.row_dimensions[header_row].height = 24
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_header
        cell.border = cell_border

    # 4. Relleno de datos de productos
    curr_row = 5
    for p in products:
        val_total = round(p['cant'] * p['precio'], 2)
        is_low = p['cant'] <= 5
        status_text = "⚠️ Stock Bajo" if is_low else "✅ Normal"

        row_vals = [
            p['id'],
            p['nombre'],
            p['cant'],
            status_text,
            p['precio'],
            val_total
        ]

        ws.row_dimensions[curr_row].height = 20
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=curr_row, column=col_idx, value=val)
            cell.border = cell_border
            
            # Formatos condicionales y estilos
            if col_idx in (1, 3):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 4:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if is_low:
                    cell.font = Font(name="Calibri", size=10, bold=True, color="DC2626")
                    cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                else:
                    cell.font = Font(name="Calibri", size=10, color="059669")
            elif col_idx == 5:
                cell.number_format = '$#,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx == 6:
                cell.number_format = '$#,##0.00'
                cell.font = Font(name="Calibri", size=11, bold=True)
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        curr_row += 1

    # 5. Fila final de totales
    ws.row_dimensions[curr_row].height = 22
    t_cell = ws.cell(row=curr_row, column=1, value="TOTALES")
    t_cell.font = Font(name="Calibri", size=11, bold=True)
    t_cell.fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    t_cell.border = double_bottom_border

    for c in range(2, 7):
        cell = ws.cell(row=curr_row, column=c)
        cell.fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
        cell.border = double_bottom_border

    stock_cell = ws.cell(row=curr_row, column=3, value=total_stock)
    stock_cell.font = Font(name="Calibri", size=11, bold=True)
    stock_cell.alignment = Alignment(horizontal="center", vertical="center")

    val_cell = ws.cell(row=curr_row, column=6, value=total_val_stock)
    val_cell.font = Font(name="Calibri", size=11, bold=True, color="059669")
    val_cell.number_format = '$#,##0.00'
    val_cell.alignment = Alignment(horizontal="right", vertical="center")

    # Auto-ajuste de anchos y guardado
    apply_auto_width(ws)
    wb.save(filepath)
    return filepath


def export_invoices_to_excel(filepath=None):
    """
    Función de conveniencia y compatibilidad que redirige a `export_invoices_by_day_and_month`.
    
    Permite generar reportes mensuales organizados por días manteniendo compatibilidad
    con llamadas legacy que suministraban un argumento opcional de ruta.

    Args:
        filepath (str, opcional): Ruta o indicador de exportación.

    Returns:
        list: Lista de rutas de los archivos Excel generados.
    """
    return export_invoices_by_day_and_month()


def export_clients_to_excel(filepath=None):
    """
    Exporta el directorio de clientes registrados a un archivo Excel formateado.
    
    Genera una tabla con las columnas: Cédula / ID, Nombre del Cliente y Teléfono.
    Aplica estilo institucional en cabeceras y autoajuste de anchos.

    Args:
        filepath (str, opcional): Ruta destino. Si es None, se almacena como
                                  'reportes/Directorio_Clientes.xlsx'.

    Returns:
        str: Ruta absoluta del archivo generado.
    """
    if filepath is None:
        reports_dir = get_reports_folder()
        filepath = os.path.join(reports_dir, "Directorio_Clientes.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Directorio de Clientes"
    
    headers = ["Cédula / ID", "Nombre del Cliente", "Teléfono de Contacto"]
    ws.append(headers)
    
    # Aplicar estilos a la fila de cabecera
    fill, font, align = create_header_style()
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = fill
        cell.font = font
        cell.alignment = align

    # Poblar filas con datos de los clientes
    clients = database.get_all_clients()
    for c in clients:
        ws.append([
            c['cedula'],
            c['nombre'],
            c['telefono'] or "N/A"
        ])
        
    apply_auto_width(ws)
    wb.save(filepath)
    return filepath
